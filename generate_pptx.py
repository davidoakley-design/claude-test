#!/usr/bin/env python3
"""
Generate a beautiful PPTX presentation with an interactive
hub-and-spoke Data Engineering Team Network diagram.

Usage:
    pip install python-pptx
    python generate_pptx.py

Output:
    data_engineering_network.pptx
"""

import math
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ── Slide dimensions (widescreen 16:9) ──
SLIDE_WIDTH = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)

# ── Center of the diagram ──
CX, CY = 6.667, 4.2

# ── Color palette ──
COLORS = {
    'bg':         RGBColor(0x0F, 0x17, 0x2A),
    'card_bg':    RGBColor(0x1E, 0x29, 0x3B),
    'text':       RGBColor(0xF1, 0xF5, 0xF9),
    'text_sec':   RGBColor(0x94, 0xA3, 0xB8),
    'core':       RGBColor(0x63, 0x66, 0xF1),
    'ml':         RGBColor(0xDB, 0x27, 0x77),
    'product':    RGBColor(0x05, 0x96, 0x69),
    'datasci':    RGBColor(0x25, 0x63, 0xEB),
    'marketing':  RGBColor(0xEA, 0x58, 0x0C),
    'platform':   RGBColor(0x7C, 0x3A, 0xED),
    'customer':   RGBColor(0x0D, 0x94, 0x88),
    'devops':     RGBColor(0xD9, 0x77, 0x06),
    'sales':      RGBColor(0xDC, 0x26, 0x26),
    'supply':     RGBColor(0x16, 0xA3, 0x4A),
    'finance':    RGBColor(0x93, 0x33, 0xEA),
    'risk':       RGBColor(0x47, 0x55, 0x69),
    'hr':         RGBColor(0xEC, 0x48, 0x99),
    'line':       RGBColor(0x33, 0x40, 0x55),
}

# ── Team data ──
teams = [
    {'id': 'core',      'name': 'Core Data\nEngineering', 'count': 22, 'color': 'core',
     'r': 0.92, 'angle': 0,   'dist': 0,
     'desc': 'Central hub for all data pipelines & infrastructure'},
    {'id': 'ml',        'name': 'ML\nEngineering',        'count': 18, 'color': 'ml',
     'r': 0.72, 'angle': 250, 'dist': 2.4,
     'desc': 'Building and deploying ML models at scale, MLOps pipelines & feature stores'},
    {'id': 'product',   'name': 'Product\nAnalytics',     'count': 15, 'color': 'product',
     'r': 0.64, 'angle': 50,  'dist': 2.5,
     'desc': 'Driving product decisions through A/B testing, funnel analysis & user behavior'},
    {'id': 'datasci',   'name': 'Data\nScience',          'count': 14, 'color': 'datasci',
     'r': 0.60, 'angle': 200, 'dist': 2.6,
     'desc': 'Advanced statistical modeling, forecasting & deep business insights'},
    {'id': 'marketing', 'name': 'Marketing\nData',        'count': 12, 'color': 'marketing',
     'r': 0.56, 'angle': 20,  'dist': 2.8,
     'desc': 'Attribution modeling, campaign analytics & marketing performance optimization'},
    {'id': 'platform',  'name': 'Platform\nEngineering',  'count': 11, 'color': 'platform',
     'r': 0.54, 'angle': 160, 'dist': 2.7,
     'desc': 'Cloud infrastructure, data platform tooling & developer experience'},
    {'id': 'customer',  'name': 'Customer\nSuccess',      'count': 10, 'color': 'customer',
     'r': 0.52, 'angle': 110, 'dist': 2.9,
     'desc': 'Customer health scoring, churn prediction & success metrics tracking'},
    {'id': 'devops',    'name': 'DevOps\nData',           'count': 9,  'color': 'devops',
     'r': 0.50, 'angle': 290, 'dist': 2.5,
     'desc': 'CI/CD for data pipelines, monitoring, observability & reliability engineering'},
    {'id': 'sales',     'name': 'Sales\nAnalytics',       'count': 8,  'color': 'sales',
     'r': 0.46, 'angle': 340, 'dist': 3.0,
     'desc': 'Revenue analytics, sales forecasting, pipeline optimization & CRM data'},
    {'id': 'supply',    'name': 'Supply\nChain',          'count': 7,  'color': 'supply',
     'r': 0.44, 'angle': 80,  'dist': 3.1,
     'desc': 'Supply chain optimization, demand forecasting & logistics analytics'},
    {'id': 'finance',   'name': 'Finance\nData',          'count': 6,  'color': 'finance',
     'r': 0.42, 'angle': 230, 'dist': 3.0,
     'desc': 'Financial reporting automation, budget analytics & revenue recognition'},
    {'id': 'risk',      'name': 'Risk &\nCompliance',     'count': 5,  'color': 'risk',
     'r': 0.38, 'angle': 315, 'dist': 2.4,
     'desc': 'Data governance, regulatory compliance, privacy engineering & risk assessment'},
    {'id': 'hr',        'name': 'HR\nAnalytics',          'count': 4,  'color': 'hr',
     'r': 0.36, 'angle': 140, 'dist': 3.0,
     'desc': 'People analytics, workforce planning & employee experience measurement'},
]


def polar_to_xy(angle_deg, dist):
    """Convert polar coordinates to x, y in inches from center."""
    rad = math.radians(angle_deg)
    return CX + dist * math.cos(rad), CY + dist * math.sin(rad)


def lighten(color, factor=0.3):
    """Lighten an RGBColor by a factor."""
    r = min(255, int(color[0] + (255 - color[0]) * factor))
    g = min(255, int(color[1] + (255 - color[1]) * factor))
    b = min(255, int(color[2] + (255 - color[2]) * factor))
    return RGBColor(r, g, b)


def add_gradient_fill(shape, color_key):
    """Add a subtle gradient fill to a shape."""
    fill = shape.fill
    fill.gradient()
    fill.gradient_stops[0].color.rgb = lighten(COLORS[color_key], 0.25)
    fill.gradient_stops[0].position = 0.0
    fill.gradient_stops[1].color.rgb = COLORS[color_key]
    fill.gradient_stops[1].position = 1.0


def center_text_vertically(shape):
    """Set vertical text anchor to center within the shape."""
    txBody = shape._element.find('.//' + qn('a:txBody'))
    if txBody is not None:
        bodyPr = txBody.find(qn('a:bodyPr'))
        if bodyPr is not None:
            bodyPr.set('anchor', 'ctr')


def draw_connector_line(slide, x1, y1, x2, y2, color=None):
    """Draw a subtle dashed connector line between two points."""
    c = color or COLORS['line']
    connector = slide.shapes.add_connector(
        1,  # straight connector
        Inches(x1), Inches(y1),
        Inches(x2), Inches(y2),
    )
    connector.begin_x = Inches(x1)
    connector.begin_y = Inches(y1)
    connector.end_x = Inches(x2)
    connector.end_y = Inches(y2)

    line = connector.line
    line.color.rgb = c
    line.width = Pt(1.2)
    line.dash_style = 4  # dash


def add_slide_hyperlink(shape, target_slide):
    """Add a click action to navigate to a specific slide using OPC relationships."""
    slide_part = shape.part  # the slide part this shape lives on
    target_part = target_slide.part

    # Create a relationship from this slide to the target slide
    rel_type = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide'
    rId = slide_part.relate_to(target_part, rel_type)

    # Add the hlinkClick element to the shape's cNvPr
    sp = shape._element
    nvSpPr = sp.find(qn('p:nvSpPr'))
    if nvSpPr is not None:
        cNvPr = nvSpPr.find(qn('p:cNvPr'))
    else:
        cNvPr = sp.find('.//' + qn('p:cNvPr'))

    if cNvPr is not None:
        hlinkClick = cNvPr.makeelement(qn('a:hlinkClick'), {
            qn('r:id'): rId,
            'action': 'ppaction://hlinksldjump',
        })
        cNvPr.append(hlinkClick)


def add_node(slide, team):
    """Add a circular node to the slide with label and count. Returns the shape."""
    if team['dist'] == 0:
        x, y = CX, CY
    else:
        x, y = polar_to_xy(team['angle'], team['dist'])

    r = team['r']
    left = Inches(x - r)
    top = Inches(y - r)
    size = Inches(r * 2)

    # Main circle
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, left, top, size, size)
    add_gradient_fill(shape, team['color'])
    shape.line.fill.background()
    shape.shadow.inherit = False

    # Text inside circle
    tf = shape.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    center_text_vertically(shape)

    lines = team['name'].split('\n')
    font_size = Pt(11) if r > 0.6 else Pt(9) if r > 0.4 else Pt(8)

    for i, line_text in enumerate(lines):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = line_text
        p.font.size = font_size
        p.font.bold = True
        p.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        p.alignment = PP_ALIGN.CENTER
        p.space_after = Pt(0)
        p.space_before = Pt(0)

    # Count line
    p = tf.add_paragraph()
    p.text = str(team['count'])
    p.font.size = Pt(10) if r > 0.6 else Pt(8)
    p.font.bold = False
    p.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    p.alignment = PP_ALIGN.CENTER
    p.space_before = Pt(2)

    return shape


def create_detail_slide(prs, team):
    """Create a detail slide for a team. Returns (slide, back_button_shape)."""
    slide_layout = prs.slide_layouts[6]  # blank
    slide = prs.slides.add_slide(slide_layout)

    # Background
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = COLORS['bg']

    # Accent bar at top
    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0), Inches(0),
        SLIDE_WIDTH, Inches(0.12)
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = COLORS[team['color']]
    bar.line.fill.background()

    # Decorative card background (added first, behind content)
    card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(3.2), Inches(0.8),
        Inches(9.2), Inches(4)
    )
    card.fill.solid()
    card.fill.fore_color.rgb = COLORS['card_bg']
    card.line.fill.background()

    # Large circle icon
    icon_size = Inches(1.8)
    icon = slide.shapes.add_shape(
        MSO_SHAPE.OVAL,
        Inches(1), Inches(1.2),
        icon_size, icon_size
    )
    add_gradient_fill(icon, team['color'])
    icon.line.fill.background()
    tf = icon.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = str(team['count'])
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    p.alignment = PP_ALIGN.CENTER
    center_text_vertically(icon)

    # Team name
    name_box = slide.shapes.add_textbox(Inches(3.5), Inches(1.2), Inches(8), Inches(1))
    tf = name_box.text_frame
    p = tf.paragraphs[0]
    p.text = team['name'].replace('\n', ' ')
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = COLORS['text']

    # Count badge
    p2 = tf.add_paragraph()
    p2.text = f'{team["count"]} team members'
    p2.font.size = Pt(18)
    p2.font.color.rgb = lighten(COLORS[team['color']], 0.4)
    p2.space_before = Pt(8)

    # Description
    desc_box = slide.shapes.add_textbox(Inches(3.5), Inches(3), Inches(8), Inches(1.5))
    tf = desc_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = team['desc']
    p.font.size = Pt(16)
    p.font.color.rgb = COLORS['text_sec']
    p.line_spacing = Pt(26)

    # "Back to Overview" button
    btn = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(0.8), Inches(6.2),
        Inches(2.2), Inches(0.6)
    )
    btn.fill.solid()
    btn.fill.fore_color.rgb = COLORS['core']
    btn.line.fill.background()
    tf = btn.text_frame
    p = tf.paragraphs[0]
    p.text = '\u2190  Back to Overview'
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    p.alignment = PP_ALIGN.CENTER
    center_text_vertically(btn)

    return slide, btn


def create_title_slide(prs):
    """Create a title/cover slide."""
    slide_layout = prs.slide_layouts[6]  # blank
    slide = prs.slides.add_slide(slide_layout)

    # Background
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = COLORS['bg']

    # Title
    title_box = slide.shapes.add_textbox(
        Inches(1), Inches(2.5), Inches(11.333), Inches(2)
    )
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = 'Data Engineering'
    p.font.size = Pt(52)
    p.font.bold = True
    p.font.color.rgb = COLORS['text']
    p.alignment = PP_ALIGN.CENTER

    p2 = tf.add_paragraph()
    p2.text = 'Team Network'
    p2.font.size = Pt(52)
    p2.font.bold = True
    p2.font.color.rgb = COLORS['core']
    p2.alignment = PP_ALIGN.CENTER

    p3 = tf.add_paragraph()
    p3.text = 'Interactive Organization Map'
    p3.font.size = Pt(18)
    p3.font.color.rgb = COLORS['text_sec']
    p3.alignment = PP_ALIGN.CENTER
    p3.space_before = Pt(16)

    # Decorative circles
    for angle, dist, size, color in [
        (30, 4.5, 0.6, 'ml'), (150, 4.0, 0.45, 'product'),
        (220, 4.8, 0.35, 'datasci'), (320, 3.8, 0.5, 'customer'),
        (80, 3.5, 0.3, 'finance'), (270, 4.2, 0.4, 'devops'),
    ]:
        x, y = polar_to_xy(angle, dist)
        y -= 0.5
        s = slide.shapes.add_shape(
            MSO_SHAPE.OVAL,
            Inches(x - size), Inches(y - size),
            Inches(size * 2), Inches(size * 2)
        )
        add_gradient_fill(s, color)
        s.line.fill.background()

    return slide


def main():
    prs = Presentation()
    prs.slide_width = SLIDE_WIDTH
    prs.slide_height = SLIDE_HEIGHT

    # ── Slide 1: Title ──
    title_slide = create_title_slide(prs)

    # ── Slide 2: Network Diagram (main overview) ──
    slide_layout = prs.slide_layouts[6]  # blank
    overview_slide = prs.slides.add_slide(slide_layout)

    # Background
    bg = overview_slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = COLORS['bg']

    # Title bar
    title_box = overview_slide.shapes.add_textbox(
        Inches(0.5), Inches(0.3), Inches(12), Inches(0.6)
    )
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = 'Data Engineering Team Network'
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = COLORS['text']

    p2 = tf.add_paragraph()
    p2.text = 'Click any team node to see details'
    p2.font.size = Pt(11)
    p2.font.color.rgb = COLORS['text_sec']
    p2.space_before = Pt(2)

    # ── Create detail slides (slides 3..N) ──
    detail_slides = {}
    back_buttons = {}
    for team in teams:
        detail_slide, back_btn = create_detail_slide(prs, team)
        detail_slides[team['id']] = detail_slide
        back_buttons[team['id']] = back_btn

    # ── Draw the diagram on the overview slide ──
    # Compute positions
    positions = {}
    for team in teams:
        if team['dist'] == 0:
            positions[team['id']] = (CX, CY)
        else:
            positions[team['id']] = polar_to_xy(team['angle'], team['dist'])

    # Draw connector lines
    core_pos = positions['core']
    for team in teams[1:]:
        pos = positions[team['id']]
        draw_connector_line(overview_slide, core_pos[0], core_pos[1], pos[0], pos[1])

    # Draw nodes and wire up hyperlinks
    for team in teams:
        node_shape = add_node(overview_slide, team)
        target_slide = detail_slides[team['id']]
        add_slide_hyperlink(node_shape, target_slide)

    # Wire up back buttons to overview slide
    for team in teams:
        add_slide_hyperlink(back_buttons[team['id']], overview_slide)

    # ── Summary slide ──
    summary_layout = prs.slide_layouts[6]
    summary_slide = prs.slides.add_slide(summary_layout)

    bg = summary_slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = COLORS['bg']

    total = sum(t['count'] for t in teams)

    title_box = summary_slide.shapes.add_textbox(
        Inches(1), Inches(1), Inches(11), Inches(1.5)
    )
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = f'{total} Data Professionals'
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = COLORS['text']
    p.alignment = PP_ALIGN.CENTER

    p2 = tf.add_paragraph()
    p2.text = f'Across {len(teams)} specialized teams'
    p2.font.size = Pt(20)
    p2.font.color.rgb = COLORS['text_sec']
    p2.alignment = PP_ALIGN.CENTER
    p2.space_before = Pt(8)

    # Team cards grid
    cols = 4
    card_w, card_h = 2.6, 1.2
    start_x = (13.333 - cols * card_w - (cols - 1) * 0.3) / 2
    start_y = 3.2

    for i, team in enumerate(teams):
        col = i % cols
        row = i // cols
        x = start_x + col * (card_w + 0.3)
        y = start_y + row * (card_h + 0.25)

        card = summary_slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(x), Inches(y),
            Inches(card_w), Inches(card_h)
        )
        card.fill.solid()
        card.fill.fore_color.rgb = COLORS['card_bg']
        card.line.fill.background()

        # Color accent dot
        dot = summary_slide.shapes.add_shape(
            MSO_SHAPE.OVAL,
            Inches(x + 0.2), Inches(y + 0.42),
            Inches(0.22), Inches(0.22)
        )
        dot.fill.solid()
        dot.fill.fore_color.rgb = COLORS[team['color']]
        dot.line.fill.background()

        # Text
        tb = summary_slide.shapes.add_textbox(
            Inches(x + 0.55), Inches(y + 0.2),
            Inches(card_w - 0.7), Inches(card_h - 0.3)
        )
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = team['name'].replace('\n', ' ')
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = COLORS['text']

        p2 = tf.add_paragraph()
        p2.text = f'{team["count"]} members'
        p2.font.size = Pt(10)
        p2.font.color.rgb = COLORS['text_sec']

    # Save
    output_path = 'data_engineering_network.pptx'
    prs.save(output_path)
    print(f'\n  Created: {output_path}')
    print(f'  Slides:  {len(prs.slides)}')
    print(f'  Teams:   {len(teams)}')
    print(f'  Total:   {total} members')
    print(f'\n  Open in PowerPoint and click any node to navigate!\n')


if __name__ == '__main__':
    main()
